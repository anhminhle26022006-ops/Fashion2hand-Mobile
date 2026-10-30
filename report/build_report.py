from pathlib import Path
import re, json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
sec.left_margin, sec.right_margin = Cm(2.7), Cm(2)
sec.footer_distance = Cm(1)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0,0,0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
for style in doc.styles:
    for border in style._element.findall('.//' + qn('w:pBdr')):
        border.getparent().remove(border)
    for rf in style._element.findall('.//' + qn('w:rFonts')):
        for attr in list(rf.attrib):
            if 'theme' in attr.lower():del rf.attrib[attr]
        for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
    if style.name.startswith('TOC'):
        style.font.size=Pt(11)
        style.paragraph_format.space_after=Pt(3)
        style.paragraph_format.line_spacing=1
doc.styles['Normal'].font.size = Pt(13)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.2
normal.space_after = Pt(7)
normal.widow_control = True
for level, size in [(1,16), (2,14), (3,13)]:
    style = doc.styles[f'Heading {level}']
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(14)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.keep_with_next = True
doc.styles['Title'].font.size = Pt(20)
doc.styles['Title'].font.bold = True
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run()
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)

def text_runs(p, text):
    # Keep identifiers intact while removing Markdown-only code delimiters.
    for i, part in enumerate(re.split(r'(`[^`]+`)', text)):
        run=p.add_run(part.strip('`') if part.startswith('`') else part)
        if part.startswith('`'):
            run.font.name='Consolas'; run.font.size=Pt(10)

def figures():
    out=ROOT/'figures';out.mkdir(exist_ok=True)
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',25)
    bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',27)
    def wrap(draw,txt,maxw,font):
        result=[];line=''
        for word in txt.split():
            test=(line+' '+word).strip()
            if draw.textlength(test,font=font)>maxw and line:result.append(line);line=word
            else:line=test
        if line:result.append(line)
        return '\n'.join(result)
    im=Image.new('RGB',(1600,860),'white');d=ImageDraw.Draw(im)
    def box(rect,txt,bolded=False):
        d.rounded_rectangle(rect,radius=12,fill='#F1F3F5',outline='#6B7280',width=2)
        f=bold if bolded else font
        text=wrap(d,txt,rect[2]-rect[0]-34,f)
        b=d.multiline_textbbox((0,0),text,font=f,spacing=9)
        x=(rect[0]+rect[2]-(b[2]-b[0]))/2;y=(rect[1]+rect[3]-(b[3]-b[1]))/2
        d.multiline_text((x,y),text,font=f,fill='black',align='center',spacing=9)
    box((555,20,1045,105),'FASHION2HAND',True)
    d.line([(800,105),(800,145),(265,145),(1335,145)],fill='#555555',width=3)
    groups=[('Tài khoản và marketplace',['Email và Google','Hồ sơ và địa chỉ','Trang chủ và tìm kiếm','Chi tiết và yêu thích','Đăng và quản lý bài']),('Giao dịch và vận chuyển',['Chat và yêu cầu FREE','Giỏ hàng và checkout','VNPAY Sandbox','Đơn hàng và GHN','Lịch sử và thông báo']),('Quản trị',['Dashboard và phân quyền','Người dùng và nội dung','Báo cáo vi phạm','Tranh chấp theo provider','Xác minh thủ công'])]
    for j,(title,items) in enumerate(groups):
        x=30+j*535;cx=x+235;d.line([(cx,145),(cx,180)],fill='#555555',width=3)
        box((x,180,x+470,270),title,True)
        for k,txt in enumerate(items):
            yy=310+k*103;box((x,yy,x+470,yy+82),txt)
    im.save(out/'phan-ra-chuc-nang.png')
    im=Image.new('RGB',(1600,1000),'white');d=ImageDraw.Draw(im)
    actors=['Android UI','Repository','Cloud Functions','VNPAY Sandbox','Firestore']
    xs=[155,475,795,1115,1435]
    for x,name in zip(xs,actors):
        d.rectangle((x-140,25,x+140,95),fill='#E8EDF2',outline='#6B7280',width=2)
        bb=d.textbbox((0,0),name,font=bold);d.text((x-(bb[2]-bb[0])/2,46),name,font=bold,fill='black')
        d.line((x,95,x,950),fill='#C0C0C0',width=2)
    steps=[(0,1,'1  Yêu cầu thanh toán theo orderId'),(1,2,'2  createVnpayPayment'),(2,4,'3  Kiểm tra đơn và lưu attempt'),(2,0,'4  Trả URL có chữ ký, mở trình duyệt'),(3,2,'5  IPN hoặc QueryDR đã xác minh'),(2,4,'6  Transaction ghi payment và receipt'),(0,2,'7  Đọc trạng thái đã xác thực'),(2,0,'8  Trả kết quả hoặc cần đối soát')]
    for k,(a,b,label) in enumerate(steps):
        y=190+k*95;x1=xs[a];x2=xs[b]
        d.line((x1,y,x2,y),fill='#333333',width=3)
        direction=1 if x2>x1 else -1
        d.polygon([(x2,y),(x2-direction*14,y-8),(x2-direction*14,y+8)],fill='#333333')
        txt=wrap(d,label,abs(x2-x1)-20,font)
        bb=d.multiline_textbbox((0,0),txt,font=font,spacing=3)
        px=(x1+x2-(bb[2]-bb[0]))/2;py=y-(bb[3]-bb[1])-16
        d.rectangle((px-5,py-3,px+bb[2]-bb[0]+5,y-10),fill='white')
        d.multiline_text((px,py),txt,font=font,fill='black',align='center',spacing=3)
    im.save(out/'tuong-tac-vnpay.png')
figures()

def figure(name,caption):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(ROOT/'figures'/name),width=Cm(16.1))
    p=doc.add_paragraph(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(11);r.italic=True

def table(rows):
    cols=len(rows[0]); t=doc.add_table(rows=1, cols=cols)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths={2:[4.3,12],3:[3.4,6.0,6.9],4:[2.2,6.2,3.6,4.3],5:[1.5,4.4,4.3,3.5,2.6]}.get(cols,[16.3/cols]*cols)
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9');borders.append(el)
    props.append(borders)
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for ci,(cell,txt) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for side,val in [('top','80'),('bottom','80'),('left','95'),('right','95')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.08
            text_runs(p,txt)
            for r in p.runs:r.font.size=Pt(10.5)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');cp.append(shade)
                for r in p.runs:r.bold=True
        trpr=t.rows[ri]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=(ROOT/'FASHION2HAND_BAO_CAO.md').read_text(encoding='utf-8').splitlines()
title=lines[0][2:]
cover_end=lines.index('# MỤC LỤC')
for s in lines[2:cover_end]:
    if not s:continue
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(13)
    text_runs(p,s)
    if s=='ĐỒ ÁN MÔN HỌC':
        p.runs[0].bold=True;p.runs[0].font.size=Pt(17)
        tp=doc.add_paragraph(title, 'Title');tp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        tp.paragraph_format.space_before=Pt(24);tp.paragraph_format.space_after=Pt(30)
doc.add_page_break()
toc_title=doc.add_paragraph('MỤC LỤC');toc_title.runs[0].bold=True;toc_title.runs[0].font.size=Pt(16)
p=doc.add_paragraph();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');p._p.append(f)
settings=doc.settings.element;up=OxmlElement('w:updateFields');up.set(qn('w:val'),'true');settings.append(up)
i=cover_end+1
code=False
while i<len(lines):
    s=lines[i];i+=1
    if not s or s.startswith('[Mục lục') or s.startswith('!['):continue
    if s.startswith('```'):
        code=not code;continue
    if code:
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1
        p.paragraph_format.keep_with_next=True
        r=p.add_run(s);r.font.name='Consolas';r.font.size=Pt(8)
        continue
    if s.startswith('|'):
        rows=[[x.strip() for x in s.strip('|').split('|')]]
        while i<len(lines) and lines[i].startswith('|'):
            row=[x.strip() for x in lines[i].strip('|').split('|')];i+=1
            if all(re.fullmatch(r':?-+:?',x) for x in row):continue
            rows.append(row)
        table(rows);continue
    if s.startswith('# '):
        h=doc.add_heading(s[2:],1);h.paragraph_format.page_break_before=True
    elif s.startswith('## '):
        doc.add_heading(s[3:],2)
        if s.startswith('## 2.2.'):
            figure('phan-ra-chuc-nang.png','Hình 2.1. Phân rã chức năng Fashion2hand theo ba nhóm nghiệp vụ.')
    elif s.startswith('### '):
        doc.add_heading(s[4:],3)
        if s.startswith('### 4.2.3.'):
            figure('tuong-tac-vnpay.png','Hình 4.1. Tương tác thanh toán rút gọn; Return không ghi nhận PAID, server xác minh qua IPN hoặc QueryDR.')
    else:
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;text_runs(p,s)

doc.core_properties.title='Báo cáo ứng dụng Fashion2hand'
doc.core_properties.subject='Mua bán và cho tặng quần áo đã qua sử dụng'
doc.core_properties.author='Nhóm Fashion2hand'
for p in doc.paragraphs:
    for border in p._p.findall('.//' + qn('w:pBdr')):border.getparent().remove(border)
out=ROOT/'FASHION2HAND_BAO_CAO.docx';doc.save(out)

# Evidence inventory omits credentials, build outputs and third-party libraries.
inventory=[]
paths=[]
for folder in ['app/src/main/java','app/src/main/res/layout','app/src/main/res/layout-land','app/src/test','app/src/androidTest','functions/test']:
    paths.extend(p for p in (REPO/folder).rglob('*') if p.is_file() and p.suffix in ['.java','.xml','.js','.ps1','.txt'])
paths.extend((REPO/'functions').glob('*.js'))
paths.extend(REPO.glob('*.md'))
paths.extend(REPO/p for p in ['firestore.rules','storage.rules','firestore.indexes.json','app/build.gradle.kts','app/src/main/AndroidManifest.xml'])
for p in sorted(set(paths)):
    raw=p.read_bytes();content=raw.decode('utf-8-sig',errors='replace')
    inventory.append({'path':p.relative_to(REPO).as_posix(),'lines':len(content.splitlines()),'sha256':hashlib.sha256(raw).hexdigest()})
(ROOT/'DOI_CHIEU_NGUON.json').write_text(json.dumps({'commit':'701a31f98ee1aae4bf69ee000a083800d1912eac','files':inventory},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(out),'words':len(' '.join(lines).split()),'source_inventory_files':len(inventory),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)},ensure_ascii=False))
